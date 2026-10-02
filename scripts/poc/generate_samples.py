#!/usr/bin/env python3
"""Generate a deterministic synthetic corpus for InterviewMirror M0a.

Requires the pinned Python runtime and generator packages in python-toolchain.lock.json.
No real personal data is used. The generated PDFs/DOCX files are benchmark fixtures.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)
from poc_runtime import check_runtime, toolchain_lock

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "poc" / "samples"
DATASET_VERSION = "interviewmirror-samples.v1.1.0"
FONT_PATH = Path("C:/Windows/Fonts/msyh.ttc")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def definitions() -> list[dict]:
    resumes = [
        ("R01", "pdf", ["text"], "Candidate-R01", "Bachelor of Computer Science", "Java, Spring Boot, RAG", "Knowledge Base Q&A", "Recall@5 0.84"),
        ("R02", "docx", ["text", "table"], "Candidate-R02", "Bachelor of Software Engineering", "Java, Vue 3, REST API", "AI Interview Practice", "P95 1.8 seconds"),
        ("R03", "pdf", ["text", "multi-page", "cross-page"], "Candidate-R03", "Bachelor of Computer Science", "Python, RAG, Evaluation", "Document Retrieval Service", "Recall@5 0.88"),
        ("R04", "docx", ["table", "multi-column"], "Candidate-R04", "Bachelor of Information Systems", "Java, Spring Boot, PostgreSQL", "Campus Knowledge Assistant", "1,200 documents"),
        ("R05", "pdf", ["scanned", "ocr", "image-only"], "Candidate-R05", "Bachelor of Computer Engineering", "Java, OCR, RAG", "Scanned Resume Parser", "F1 0.91"),
        ("R06", "docx", ["text", "bullets"], "Candidate-R06", "Bachelor of Artificial Intelligence", "Python, LangGraph, Spring AI", "Interview Report Agent", "40 test cases"),
        ("R07", "pdf", ["table", "formula"], "Candidate-R07", "Bachelor of Computer Science", "Java, Redis, Vue 3", "Model Gateway", "P95 = 420 ms"),
        ("R08", "docx", ["multi-column", "multi-page", "cross-page"], "Candidate-R08", "Bachelor of Data Science", "Python, SQL, RAG", "Question Bank Builder", "Question recall 0.93"),
        ("R09", "pdf", ["text", "formula"], "Candidate-R09", "Bachelor of Software Engineering", "Java, Spring Boot, SSE", "Streaming Answer UI", "Retry rate 1.2%"),
        ("R10", "docx", ["table", "cross-page"], "Candidate-R10", "Bachelor of Computer Science", "TypeScript, Java, Docker", "Local AI Demo", "Startup 12 minutes"),
    ]
    banks = [
        ("Q01", "md", ["text"], ["Explain the RAG retrieval pipeline.", "How would you build an offline evaluation set?", "How do you measure citation correctness?"]),
        ("Q02", "txt", ["text", "numbered"], ["Describe a project you owned end to end.", "How did you choose the database schema?", "What trade-off did you make under a deadline?"]),
        ("Q03", "md", ["table", "formula"], ["How do you calculate Recall@K?", "When would you use reranking?", "How do you compare two prompt versions?"]),
        ("Q04", "docx", ["office", "table"], ["Design a resilient model API client.", "How should retries use exponential backoff?", "How do you prevent duplicate report creation?"]),
        ("Q05", "pdf", ["two-column", "table", "complex"], ["Explain SSE reconnect behavior.", "How do you persist partial model output?", "What belongs in a trace identifier?"]),
        ("Q06", "txt", ["multiline", "numbered"], ["How would you design resume ownership checks?", "What should be deleted with a resume?", "How do you keep secrets out of logs?"]),
        ("Q07", "md", ["multi-page", "cross-page"], ["Compare polling with server-sent events.", "How do you set a report generation timeout?", "How would you test an early interview end?"]),
        ("Q08", "pdf", ["scanned", "ocr", "image-only"], ["How do you detect an OCR extraction failure?", "How can a user correct parsed questions?", "When should a document be marked unprocessable?"]),
        ("Q09", "md", ["table", "multi-page", "complex"], ["How do you prioritize a product gap?", "What makes evidence trustworthy?", "How do you mark an unassessed skill?"]),
        ("Q10", "txt", ["text", "bullets"], ["Describe a production incident you investigated.", "What metric proved your fix worked?", "What would you improve in the next iteration?"]),
    ]
    jds = [
        ("J01", "docx", ["office", "table", "complex"], "AI application intern", ["Build enterprise knowledge assistants using RAG.", "Create offline evaluation datasets and improve answer quality.", "Handle API timeouts, retries and fallbacks."]),
        ("J02", "pdf", ["office", "two-column", "formula", "complex"], "AI full-stack intern", ["Deliver Vue and Java features end to end.", "Stream model output with SSE and preserve session state.", "Keep API p95 below 800 ms for non-model endpoints."]),
        ("J03", "docx", ["office", "formula", "cross-page", "complex"], "AI platform intern", ["Track model usage and per-session cost.", "Keep monthly model budget below CNY 300.", "Provide evidence-backed interview reports."]),
        ("J04", "pdf", ["scanned", "ocr", "image-only"], "RAG engineering intern", ["Implement document parsing and citation locators.", "Evaluate scanned PDF OCR quality.", "Collaborate through written design reviews."]),
    ]
    result = []
    for sid, fmt, tags, name, education, skills, project, metric in resumes:
        result.append({
            "sampleId": sid, "category": "resume", "format": fmt, "complexityTags": tags,
            "title": f"Synthetic Resume {sid}", "fields": {
                "name": name, "education": education, "skills": skills,
                "project": project, "metric": metric,
            },
        })
    for sid, fmt, tags, questions in banks:
        result.append({
            "sampleId": sid, "category": "question_bank", "format": fmt,
            "complexityTags": tags, "title": f"Synthetic Question Bank {sid}",
            "questions": questions,
        })
    for sid, fmt, tags, title, requirements in jds:
        result.append({
            "sampleId": sid, "category": "jd", "format": fmt,
            "complexityTags": tags, "title": title, "requirements": requirements,
        })
    return result


def text_for(sample: dict) -> str:
    lines = [sample["title"], "Source: synthetic benchmark fixture; no real personal data."]
    if sample["category"] == "resume":
        labels = {"name": "Name", "education": "Education", "skills": "Skills", "project": "Project", "metric": "Metric"}
        lines.extend(f"{labels[k]}: {v}" for k, v in sample["fields"].items())
        lines.extend(["Experience: Built and evaluated a small local AI application.", "Contact: synthetic@example.invalid"])
    elif sample["category"] == "question_bank":
        lines.extend(f"Q{i}: {q}" for i, q in enumerate(sample["questions"], 1))
        if "formula" in sample["complexityTags"]:
            lines.append("Formula: Recall@K = relevant retrieved items / all relevant items")
    else:
        lines.extend([f"Role: {sample['title']}"])
        lines.extend(f"Requirement {i}: {req}" for i, req in enumerate(sample["requirements"], 1))
        if "formula" in sample["complexityTags"]:
            lines.append("Constraint: p95 = the 95th percentile of request latency")
    return "\n".join(lines)


def register_pdf_font() -> str:
    pdfmetrics.registerFont(TTFont("InterviewMirrorCJK", str(FONT_PATH), subfontIndex=0))
    return "InterviewMirrorCJK"


def render_scanned_pdf(sample: dict, path: Path) -> None:
    font = ImageFont.truetype(str(FONT_PATH), 31)
    image = Image.new("RGB", (1275, 1800), "white")
    draw = ImageDraw.Draw(image)
    y = 95
    for raw in text_for(sample).splitlines():
        # OCR-friendly line lengths, with generous line spacing and dark text.
        wrapped = [raw[i:i + 47] for i in range(0, len(raw), 47)] or [""]
        for line in wrapped:
            draw.text((96, y), line, fill="#111111", font=font)
            y += 58
    image.save(path.with_suffix(".scan.png"), dpi=(180, 180))
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle(sample["title"])
    c.setAuthor("InterviewMirror synthetic corpus")
    c.drawImage(str(path.with_suffix(".scan.png")), 0, 0, width=A4[0], height=A4[1])
    c.save()
    path.with_suffix(".scan.png").unlink(missing_ok=True)


def render_pdf(sample: dict, path: Path) -> None:
    if "scanned" in sample["complexityTags"]:
        render_scanned_pdf(sample, path)
        return
    font_name = register_pdf_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("FixtureBody", parent=styles["BodyText"], fontName=font_name,
                          fontSize=10, leading=16, alignment=TA_LEFT, spaceAfter=7)
    heading = ParagraphStyle("FixtureTitle", parent=styles["Title"], fontName=font_name,
                             fontSize=16, leading=21, spaceAfter=12)
    story = [Paragraph(sample["title"], heading), Paragraph("Synthetic fixture · no real personal data", body)]
    if sample["category"] == "resume" and "table" in sample["complexityTags"]:
        rows = [["Field", "Value"]] + [[k, v] for k, v in sample["fields"].items()]
        table = Table(rows, colWidths=[38 * mm, 118 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f2d677")),
            ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#a7a7a7")),
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("PADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([table, Spacer(1, 8)])
    if "two-column" in sample["complexityTags"]:
        if sample["category"] == "question_bank":
            cells = [[Paragraph("<br/>".join(sample["questions"][:2]), body),
                      Paragraph(sample["questions"][2], body)]]
        else:
            cells = [[Paragraph("Name / Education / Skills<br/>" + text_for(sample).replace("\n", "<br/>"), body),
                      Paragraph("Column B<br/>Experience and project notes", body)]]
        two_col = Table(cells, colWidths=[78 * mm, 78 * mm])
        two_col.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                     ("BOX", (0, 0), (-1, -1), .5, colors.grey),
                                     ("INNERGRID", (0, 0), (-1, -1), .3, colors.lightgrey)]))
        story.extend([two_col, Spacer(1, 8)])
    if sample["category"] == "question_bank" and "table" in sample["complexityTags"]:
        story.append(Table([["#", "Question"], *[[str(i), q] for i, q in enumerate(sample["questions"], 1)]],
                           colWidths=[12 * mm, 140 * mm], repeatRows=1,
                           style=TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.grey),
                                             ("FONTNAME", (0, 0), (-1, -1), font_name)])))
    elif sample["category"] == "jd" and "table" in sample["complexityTags"]:
        story.append(Table([["Importance", "Requirement"], *[["Core" if i < 3 else "Preferred", r]
                               for i, r in enumerate(sample["requirements"], 1)]],
                           colWidths=[30 * mm, 122 * mm], repeatRows=1,
                           style=TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.grey),
                                             ("FONTNAME", (0, 0), (-1, -1), font_name)])))
    else:
        for line in text_for(sample).splitlines()[1:]:
            story.append(Paragraph(line, body))
    if "multi-page" in sample["complexityTags"] or "cross-page" in sample["complexityTags"]:
        story.extend([PageBreak(), Paragraph("Continuation · " + sample["sampleId"], heading)])
        for i in range(1, 7):
            story.append(Paragraph(f"Continuation item {i}: {text_for(sample).splitlines()[2]}", body))
    class InvariantCanvas(__import__("reportlab.pdfgen.canvas", fromlist=["Canvas"]).Canvas):
        def __init__(self, *args, **kwargs):
            kwargs["invariant"] = 1
            super().__init__(*args, **kwargs)
            self.setTitle(sample["title"])
            self.setAuthor("InterviewMirror synthetic corpus")

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=18 * mm, bottomMargin=18 * mm, pageCompression=0).build(story, canvasmaker=InvariantCanvas)


def render_docx(sample: dict, path: Path) -> None:
    doc = Document()
    fixed_time = __import__("datetime").datetime(2000, 1, 1, 0, 0, 0)
    doc.core_properties.created = fixed_time
    doc.core_properties.modified = fixed_time
    doc.core_properties.last_modified_by = "InterviewMirror synthetic corpus"
    doc.core_properties.revision = 1
    doc.add_heading(sample["title"], level=1)
    doc.add_paragraph("Synthetic fixture · no real personal data")
    if sample["category"] == "resume":
        fields = sample["fields"]
        if "table" in sample["complexityTags"]:
            table = doc.add_table(rows=0, cols=2)
            table.style = "Light Shading Accent 1"
            for key, value in fields.items():
                cells = table.add_row().cells
                cells[0].text = key.title()
                cells[1].text = value
        else:
            for key, value in fields.items():
                doc.add_paragraph(f"{key.title()}: {value}")
        if "multi-column" in sample["complexityTags"]:
            doc.add_paragraph("Column A: " + fields["project"])
            doc.add_paragraph("Column B: " + fields["metric"])
        if "cross-page" in sample["complexityTags"]:
            doc.add_page_break()
            doc.add_heading("Additional experience", level=2)
            doc.add_paragraph("Experience: Local interview report generator")
    elif sample["category"] == "question_bank":
        if "table" in sample["complexityTags"]:
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Shading Accent 1"
            table.rows[0].cells[0].text = "ID"
            table.rows[0].cells[1].text = "Question"
            for i, question in enumerate(sample["questions"], 1):
                cells = table.add_row().cells
                cells[0].text = f"Q{i}"
                cells[1].text = question
        else:
            for i, question in enumerate(sample["questions"], 1):
                doc.add_paragraph(f"Q{i}: {question}")
    else:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Shading Accent 1"
        table.rows[0].cells[0].text = "Importance"
        table.rows[0].cells[1].text = "Requirement"
        for i, req in enumerate(sample["requirements"], 1):
            cells = table.add_row().cells
            cells[0].text = "Core" if i < 3 else "Preferred"
            cells[1].text = req
        if "formula" in sample["complexityTags"]:
            doc.add_paragraph("Constraint: Recall@5 = relevant retrieved items / all relevant items")
        if "cross-page" in sample["complexityTags"]:
            doc.add_page_break()
            doc.add_paragraph("Continuation: " + sample["requirements"][-1])
    doc.save(path)


def normalize_docx_zip(path: Path) -> None:
    """Remove ZIP member timestamp variance so corpus hashes reproduce byte-for-byte."""
    source = path.read_bytes()
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source), "r") as original, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as normalized:
        for info in sorted(original.infolist(), key=lambda item: item.filename):
            out = zipfile.ZipInfo(info.filename, date_time=(2000, 1, 1, 0, 0, 0))
            out.compress_type = zipfile.ZIP_DEFLATED
            out.external_attr = info.external_attr
            out.create_system = info.create_system
            out.flag_bits = info.flag_bits
            normalized.writestr(out, original.read(info.filename))
    path.write_bytes(buffer.getvalue())


def render(sample: dict, path: Path) -> None:
    if sample["format"] == "docx":
        render_docx(sample, path)
    elif sample["format"] == "pdf":
        render_pdf(sample, path)
    else:
        content = text_for(sample)
        if sample["format"] == "md":
            content = "# " + sample["title"] + "\n\n" + "\n\n".join(
                f"## Q{i}: {q}" for i, q in enumerate(sample.get("questions", []), 1)
            ) if sample["category"] == "question_bank" else "# " + sample["title"] + "\n\n" + content
        path.write_text(content + "\n", encoding="utf-8")


def main() -> None:
    check_runtime()
    lock = toolchain_lock()
    OUT.mkdir(parents=True, exist_ok=True)
    # Remove only files generated by this script, never arbitrary corpus files.
    for old in OUT.glob("*.manifest.json"):
        old.unlink()
    entries = []
    for sample in definitions():
        extension = sample["format"]
        path = OUT / f"{sample['sampleId']}.{extension}"
        render(sample, path)
        if extension == "docx":
            normalize_docx_zip(path)
        record = {
            "schemaVersion": DATASET_VERSION,
            "sampleId": sample["sampleId"],
            "file": path.name,
            "format": sample["format"],
            "mediaType": {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                          "md": "text/markdown", "txt": "text/plain"}[extension],
            "category": sample["category"],
            "complexityTags": sample["complexityTags"],
            "source": "synthetic; generated locally by scripts/poc/generate_samples.py",
            "scanned": "scanned" in sample["complexityTags"],
            "groundTruth": {k: v for k, v in sample.items() if k in ("fields", "questions", "requirements")},
            "sha256": sha256(path),
        }
        (OUT / f"{sample['sampleId']}.manifest.json").write_text(
            json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        entries.append(record)
    global_manifest = {
        "schemaVersion": DATASET_VERSION,
        "datasetId": "interviewmirror-stage0-synthetic-24",
        "sampleCount": len(entries),
        "generatedBy": "scripts/poc/generate_samples.py",
        "generator": {"python": lock["python"], "platform": lock["platform"],
                      "packages": lock["packages"], "systemInputs": lock["systemInputs"]},
        "samples": [{"sampleId": e["sampleId"], "file": e["file"], "sha256": e["sha256"]} for e in entries],
    }
    (OUT / "manifest.json").write_text(json.dumps(global_manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"sampleCount": len(entries), "output": str(OUT), "datasetSha256": sha256(OUT / "manifest.json")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
