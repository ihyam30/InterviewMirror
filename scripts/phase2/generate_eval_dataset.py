#!/usr/bin/env python3
"""Build the Stage 2 deterministic 30-document evaluation set.

The original 20 Stage 0 resume/question-bank fixtures are referenced by hash;
only the additional 10 fixtures are written into the Stage 2 directory.
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
POC = ROOT / "scripts" / "poc"
sys.path.insert(0, str(POC))
from poc_runtime import check_runtime, generator_metadata, toolchain_lock  # noqa: E402
from generate_samples import normalize_docx_zip  # noqa: E402

OUT = ROOT / "data" / "poc" / "phase2" / "samples"
OLD = ROOT / "data" / "poc" / "samples"
DATASET = ROOT / "data" / "poc" / "phase2" / "manifest.json"
ANNOTATIONS = ROOT / "data" / "poc" / "phase2" / "annotations"


def sha(path: Path) -> str:
    data = path.read_bytes()
    if path.suffix.lower() in {".json", ".md", ".txt"}:
        data = data.replace(b"\r\n", b"\n").replace(b"\r", b"\n").replace(b"\n", b"\r\n")
    return hashlib.sha256(data).hexdigest()


def resume_fixture(number: int) -> tuple[dict, str, str]:
    sample_id = f"R{number:02d}"
    spec = {
        "sampleId": sample_id,
        "format": ["pdf", "docx", "pdf", "docx", "pdf"][number - 11],
        "complexityTags": [
            ["table", "multi-page"], ["table", "office"], ["multi-column"],
            ["bilingual", "multi-page"], ["scanned", "ocr"],
        ][number - 11],
        "fields": {
            "name": f"Candidate-{sample_id}",
            "education": ["Bachelor of Computer Science", "Bachelor of Software Engineering",
                          "Master of Information Systems", "Bachelor of Artificial Intelligence",
                          "Bachelor of Computer Engineering"][number - 11],
            "skills": ["Java, Spring Boot, PostgreSQL", "Python, Vue 3, RAG",
                       "Java, LangGraph4j, Docker", "Python, SQL, Evaluation",
                       "Java, OCR, Document Parsing"][number - 11],
            "project": ["Local Interview Practice", "Question Bank Importer", "Model Gateway",
                        "Resume Parsing Pipeline", "Synthetic OCR Evaluation"][number - 11],
            "metric": ["Parse p95 38 seconds", "Recall@5 0.91", "P95 420 ms",
                       "F1 0.94", "Question recall 0.92"][number - 11],
        },
    }
    lines = [
        f"Synthetic Resume {sample_id}",
        "Source: synthetic benchmark fixture; no real personal data.",
        f"Name: {spec['fields']['name']}",
        f"Education: {spec['fields']['education']}",
        f"Skills: {spec['fields']['skills']}",
        f"Project: {spec['fields']['project']}",
        f"Metric: {spec['fields']['metric']}",
        "Experience: Built and evaluated a small local AI application.",
    ]
    return spec, "\n".join(lines), f"Synthetic Resume {sample_id}"


def bank_fixture(number: int) -> tuple[dict, str, str]:
    sample_id = f"Q{number:02d}"
    questions = {
        11: ["How would you design a reliable document parsing task?", "How do you prevent duplicate worker claims?", "How should an expired task lease be recovered?"],
        12: ["How do you validate uploaded PDF and DOCX files?", "What data should a parse task persist?", "How do you handle deletion when object storage fails?"],
        13: ["How would you evaluate resume field extraction quality?", "How should question matching avoid duplicate true positives?", "What evidence belongs in a parse result?"],
        14: ["How do you keep confirmed resume data trustworthy?", "What should happen after editing confirmed content?", "How do you restrict document access by owner?"],
        15: ["How would you recover after the backend restarts?", "How do you measure parsing p95 latency?", "How should parser errors be shown to users?"],
    }[number]
    spec = {"sampleId": sample_id, "format": ["pdf", "docx", "md", "pdf", "txt"][number - 11],
            "complexityTags": [["two-column", "table", "complex"], ["office", "table"], ["markdown", "headings"], ["table", "multi-page"], ["text", "numbered"]][number - 11],
            "questions": questions}
    text = "# Question Bank " + sample_id + "\n\n" + "\n".join(f"Q{i}: {q}" for i, q in enumerate(questions, 1))
    return spec, text, "Question Bank " + sample_id


def render_docx(path: Path, title: str, body: str, category: str) -> None:
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph("Synthetic fixture; no real personal data.")
    if category == "question_bank":
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Question"
        table.rows[0].cells[1].text = "Answer"
        for line in body.splitlines():
            if line.startswith("Q") and ": " in line:
                row = table.add_row().cells
                row[0].text = line.split(": ", 1)[1]
                row[1].text = "Synthetic reference answer"
    else:
        for line in body.splitlines()[2:]:
            label, value = line.split(": ", 1)
            row = doc.add_table(rows=1, cols=2).rows[0].cells
            row[0].text = label
            row[1].text = value
    doc.save(path)
    normalize_docx_zip(path)


def render_pdf(path: Path, title: str, body: str, category: str, complex_layout: bool = False,
               scanned: bool = False) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    if font_path.exists():
        pdfmetrics.registerFont(TTFont("EvalFont", str(font_path)))
        font_name = "EvalFont"
    else:
        font_name = "Helvetica"
    if scanned:
        from PIL import Image, ImageDraw, ImageFont
        image_path = path.with_suffix(".png")
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 32)
        image = Image.new("RGB", (1275, 1800), "white")
        draw = ImageDraw.Draw(image)
        y = 100
        for line in [title, *body.splitlines()]:
            draw.text((90, y), line, fill="#111111", font=font)
            y += 78
        image.save(image_path, dpi=(180, 180))
        pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1)
        pdf.drawImage(str(image_path), 0, 0, width=A4[0], height=A4[1])
        pdf.save()
        image_path.unlink(missing_ok=True)
        return
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle("Eval", parent=styles["BodyText"], fontName=font_name, fontSize=10, leading=15)
    story = [Paragraph(title, styles["Title"]), Spacer(1, 6)]
    if category == "question_bank":
        rows = [["#", "Question", "Answer"]]
        for index, line in enumerate([x for x in body.splitlines() if x.startswith("Q")], 1):
            rows.append([str(index), Paragraph(line.split(": ", 1)[1], body_style), "Synthetic answer"])
        if complex_layout:
            left = [Paragraph("<br/>".join([r[1].getPlainText() if hasattr(r[1], "getPlainText") else str(r[1]) for r in rows[1:2]]), body_style)]
            right = [Paragraph("<br/>".join([r[1].getPlainText() if hasattr(r[1], "getPlainText") else str(r[1]) for r in rows[2:]]), body_style)]
            story.append(Table([[left[0], right[0]]], colWidths=[82 * mm, 82 * mm], style=TableStyle([("VALIGN", (0,0),(-1,-1),"TOP"), ("BOX",(0,0),(-1,-1),.5,colors.grey), ("INNERGRID",(0,0),(-1,-1),.3,colors.lightgrey)])))
            story.append(Spacer(1, 8))
        story.append(Table(rows, colWidths=[12*mm, 115*mm, 35*mm], repeatRows=1,
                           style=TableStyle([("GRID", (0,0),(-1,-1),.4,colors.grey), ("FONTNAME",(0,0),(-1,-1),font_name), ("VALIGN",(0,0),(-1,-1),"TOP")])) )
    else:
        values = [line.split(": ", 1) for line in body.splitlines() if ": " in line]
        if complex_layout:
            rows = [["Field", "Value"], *values]
            story.append(Table(rows, colWidths=[38*mm, 120*mm], repeatRows=1,
                               style=TableStyle([("GRID",(0,0),(-1,-1),.4,colors.grey), ("FONTNAME",(0,0),(-1,-1),font_name)])))
        else:
            for key, value in values:
                story.append(Paragraph(f"<b>{key}:</b> {value}", body_style))
                story.append(Spacer(1, 5))
    if "multi-page" in body:
        story.append(PageBreak())
        story.append(Paragraph("Continuation · synthetic content", body_style))
    class InvariantCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            kwargs["invariant"] = 1
            super().__init__(*args, **kwargs)
            self.setTitle(title)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18*mm, rightMargin=18*mm,
                      topMargin=18*mm, bottomMargin=18*mm).build(story, canvasmaker=InvariantCanvas)


def main() -> None:
    check_runtime()
    OUT.mkdir(parents=True, exist_ok=True)
    ANNOTATIONS.mkdir(parents=True, exist_ok=True)
    samples = []
    for number in range(1, 11):
        for sample_id, category in ((f"R{number:02d}", "resume"), (f"Q{number:02d}", "question_bank")):
            manifest = json.loads((OLD / f"{sample_id}.manifest.json").read_text(encoding="utf-8"))
            source_path = OLD / manifest["file"]
            source_hash = sha(source_path)
            if source_hash != manifest["sha256"]:
                raise SystemExit(f"Stage 0 source hash mismatch: {sample_id}")
            samples.append({"sampleId": sample_id, "category": category, "format": manifest["format"],
                            "complexityTags": manifest["complexityTags"], "file": f"samples/{manifest['file']}",
                            "source": "Stage 0 synthetic fixture, hash-verified", "sha256": source_hash,
                            "groundTruth": manifest["groundTruth"]})
    for number in range(11, 16):
        for spec, body, title in (resume_fixture(number), bank_fixture(number)):
            sample_id = spec["sampleId"]
            category = "resume" if sample_id.startswith("R") else "question_bank"
            path = OUT / f"{sample_id}.{spec['format']}"
            if spec["format"] in {"md", "txt"}:
                path.write_text(body + "\n", encoding="utf-8", newline="\n")
            elif spec["format"] == "docx":
                render_docx(path, title, body, category)
            elif spec["format"] == "pdf":
                render_pdf(path, title, body, category, "table" in spec["complexityTags"], "scanned" in spec["complexityTags"])
            gt = {"fields": spec["fields"]} if category == "resume" else {"questions": spec["questions"]}
            sample = {"sampleId": sample_id, "category": category, "format": spec["format"],
                      "complexityTags": spec["complexityTags"], "file": f"phase2/samples/{path.name}",
                      "source": "Synthetic fixture authored for Stage 2; no personal data",
                      "sha256": sha(path), "groundTruth": gt}
            samples.append(sample)
    samples.sort(key=lambda row: row["sampleId"])
    for sample in samples:
        annotation = {"schemaVersion": "interviewmirror.phase2-annotation.v1", "sampleId": sample["sampleId"],
                      "category": sample["category"], "sha256": sample["sha256"], "groundTruth": sample["groundTruth"],
                      "annotationSource": sample["source"]}
        (ANNOTATIONS / f"{sample['sampleId']}.json").write_text(json.dumps(annotation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    dataset = {"schemaVersion": "interviewmirror.phase2-eval.v1", "datasetId": "interviewmirror-stage2-synthetic-30",
               "sampleCount": len(samples), "categoryCounts": {"resume": 15, "question_bank": 15},
               "generator": generator_metadata(toolchain_lock()), "samples": samples}
    DATASET.write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"dataset={dataset['datasetId']} samples={len(samples)} resumes=15 questionBanks=15")
    print(f"manifestSha256={sha(DATASET)} path={DATASET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
